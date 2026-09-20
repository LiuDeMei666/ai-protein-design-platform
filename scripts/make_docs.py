#!/usr/bin/env python
"""把 Markdown 交付文档转换为排版规范的 Word（.docx）文件。

为什么不用 docx-js
------------------
docx 技能默认推荐 docx-js，但本机 Node 为 v12.22.9，无法运行其所需的
现代 Node 运行时。因此改用 ``python-docx`` 实现同样的排版目标：
标题层级、正文、无序/有序列表、表格（含表头底纹与边框）、代码块（等宽字体）、
引用块。转换结果与 Markdown 源文档内容一致。

用法::

    python scripts/make_docs.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / "docs"

FONT_CN = "微软雅黑"
FONT_MONO = "Consolas"

#: 待转换的文档：(Markdown 文件名, Word 文件名)
DOCUMENTS: tuple[tuple[str, str], ...] = (
    ("部署文档.md", "部署文档.docx"),
    ("平台使用手册.md", "平台使用手册.docx"),
    ("算法原理说明.md", "算法原理说明.docx"),
    ("验证报告_胶原蛋白与工业酶.md", "验证报告_胶原蛋白与工业酶.docx"),
)

INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_CODE = re.compile(r"`([^`]+)`")
INLINE_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _set_cell_shading(cell, color: str) -> None:
    """给表格单元格加底纹。"""
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _configure_styles(document: Document) -> None:
    """设置正文与标题字体（中英文都要指定，否则中文会回退到默认字体）。"""
    normal = document.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 11)):
        style = document.styles[name]
        style.font.name = FONT_CN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def _add_runs(paragraph, text: str) -> None:
    """把内联 Markdown（粗体 / 行内代码 / 链接）写入段落。"""
    # 先处理链接，把 [text](url) 变成 "text（url）"，避免 URL 丢失
    text = INLINE_LINK.sub(lambda match: f"{match.group(1)}（{match.group(2)}）", text)

    position = 0
    pattern = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        if match.group(1) is not None:
            run = paragraph.add_run(match.group(1))
            run.bold = True
        else:
            run = paragraph.add_run(match.group(2))
            run.font.name = FONT_MONO
            run.font.size = Pt(9.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_MONO)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _add_table(document: Document, rows: list[list[str]]) -> None:
    """写入 Markdown 管道表。"""
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index in range(columns):
            cell = cells[column_index]
            value = row[column_index] if column_index < len(row) else ""
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            _add_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
            if row_index == 0:
                _set_cell_shading(cell, "D9E2F3")
                for run in paragraph.runs:
                    run.font.bold = True
    document.add_paragraph()


def _add_code_block(document: Document, lines: list[str]) -> None:
    """写入代码块（等宽、缩进、浅灰底）。"""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15

    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), "F2F2F2")
    properties.append(shading)

    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = FONT_MONO
        run.font.size = Pt(9)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_MONO)
        if index < len(lines) - 1:
            run.add_break()


def convert(markdown_path: Path, output_path: Path) -> dict:
    """转换单个 Markdown 文件。"""
    text = markdown_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    document = Document()
    _configure_styles(document)

    stats = {"headings": 0, "tables": 0, "code_blocks": 0, "lists": 0, "paragraphs": 0}

    index = 0
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    title_done = False

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            _add_table(document, table_rows)
            stats["tables"] += 1
            table_rows = []

    while index < len(lines):
        line = lines[index]

        # 代码块
        if line.strip().startswith("```"):
            if in_code:
                _add_code_block(document, code_lines)
                stats["code_blocks"] += 1
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        # 表格
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell):
                index += 1
                continue  # 分隔行
            table_rows.append(cells)
            index += 1
            continue
        flush_table()

        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        # 水平线
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            properties = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BFBFBF")
            borders.append(bottom)
            properties.append(borders)
            index += 1
            continue

        # 标题
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            content = heading_match.group(2)
            if level == 1 and not title_done:
                paragraph = document.add_heading(level=0)
                _add_runs(paragraph, content)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_done = True
            else:
                paragraph = document.add_heading(level=min(level, 4))
                _add_runs(paragraph, content)
            stats["headings"] += 1
            index += 1
            continue

        # 引用
        if stripped.startswith(">"):
            content = stripped.lstrip(">").strip()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            _add_runs(paragraph, content)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            index += 1
            continue

        # 无序列表
        bullet_match = re.match(r"^[-*+]\s+(.*)$", stripped)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, bullet_match.group(1))
            stats["lists"] += 1
            index += 1
            continue

        # 有序列表
        ordered_match = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if ordered_match:
            paragraph = document.add_paragraph(style="List Number")
            _add_runs(paragraph, ordered_match.group(1))
            stats["lists"] += 1
            index += 1
            continue

        # 普通段落
        paragraph = document.add_paragraph()
        _add_runs(paragraph, stripped)
        stats["paragraphs"] += 1
        index += 1

    flush_table()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    stats["output"] = str(output_path)
    return stats


def main() -> int:
    print("=" * 74)
    print(" 生成 Word 交付文档")
    print("=" * 74)

    failed: list[str] = []
    for markdown_name, docx_name in DOCUMENTS:
        source = DOCS_DIR / markdown_name
        if not source.exists():
            print(f"  [跳过] {markdown_name} 不存在")
            failed.append(markdown_name)
            continue
        target = DOCS_DIR / docx_name
        stats = convert(source, target)
        size_kb = target.stat().st_size / 1024
        print(
            f"  [完成] {docx_name:34s} {size_kb:7.0f} KB · "
            f"标题 {stats['headings']} · 表格 {stats['tables']} · "
            f"代码块 {stats['code_blocks']} · 列表 {stats['lists']}"
        )

    print("=" * 74)
    if failed:
        print(f" 有 {len(failed)} 个文档未生成")
        return 1
    print(" 全部 Word 文档已生成")
    print("=" * 74)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
